#!/usr/bin/env python3
"""Word-Export der Outreach-Unterlagen — liest AUSSCHLIESSLICH aus leads.json.

Erzeugt für jeden TIER1-Lead (alle Pipelines) ein vollständiges Anschreiben
(Betreff + personalisierter Brieftext) plus eine Quick-Reference-Tabelle.
Vorlauf: Validierungs-Gate.
"""
import json
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

PIPELINE_INTRO = {
    "IoT": ("wir entwickeln seit über 20 Jahren IoT- und Anlagenverwaltungs-Plattformen — "
            "unter anderem 24 Jahre lang die komplette Betriebsführung für den "
            "Energieversorger HanseWerk Natur."),
    "Energie": ("wir haben 24 Jahre lang die komplette Betriebs- und Anlagenverwaltung für "
                "HanseWerk Natur entwickelt und betrieben — Architektur, Entwicklung, "
                "Testing und Betrieb aus einer Hand."),
    "Energie_ETRM": ("wir entwickeln Software für Energiehandel und Risikomanagement und kennen "
                     "die Anforderungen physischer Händler aus über 20 Jahren Energiebranche."),
    "Telematik": ("wir entwickeln seit über 20 Jahren Embedded- und Telematik-Software — von der "
                  "Steuergeräte-Ebene bis zur Cloud-Plattform."),
    "TISAX": ("wir sind Ihr TISAX-konformer Partner für IT-Entwicklung und Cloud-Betrieb "
              "(IoT, KI, Datenplattformen) — speziell für Automotive-Zulieferer."),
}
CTA = ("Hätten Sie 20 Minuten für einen kurzen Austausch? Gerne schicke ich Ihnen vorab "
       "ein 2-minütiges Video, das unsere Plattform zeigt.")


def variant(lead):
    topic = ((lead.get("topic_relevance") or {}).get("topic_affinity") or {}).get("topic", "")
    if lead.get("pipeline") == "Energie" and "ETRM" in topic:
        return "Energie_ETRM"
    return lead.get("pipeline")


def salutation(contacts):
    for c in contacts or []:
        nm = (c.get("name") or "").strip()
        if nm and not nm.startswith("(") and "vakant" not in nm.lower():
            last = nm.split(",")[0].split()[-1]
            return f"Sehr geehrte/r Frau/Herr {last},", nm
    return "Sehr geehrte Damen und Herren,", None


def main():
    gate = subprocess.run([sys.executable, "scripts/validate_leads.py"],
                          capture_output=True, text=True)
    if gate.returncode != 0:
        print(gate.stdout)
        sys.exit("ABBRUCH: Validierungs-Gate fehlgeschlagen.")

    leads = json.load(open("leads.json"))["leads"]
    tier1 = sorted(
        [l for l in leads if l.get("tier") == "TIER1" and l.get("outreach")],
        key=lambda x: (x.get("pipeline") or "", -(x.get("overall_score") or 0)),
    )

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_heading("Direktansprache — TIER 1 Leads", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"{len(tier1)} vollständige Anschreiben über alle Pipelines · "
                    "[VIDEO-LINK] vor Versand ersetzen")
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(120, 120, 120)
    doc.add_page_break()

    current_pipeline = None
    for lead in tier1:
        co = lead["company"]
        v = variant(lead)
        intro = PIPELINE_INTRO.get(v, PIPELINE_INTRO.get(lead.get("pipeline"), ""))
        outreach = lead["outreach"]
        sal, contact_name = salutation(lead.get("contact_persons"))

        if lead.get("pipeline") != current_pipeline:
            current_pipeline = lead.get("pipeline")
            h = doc.add_heading(f"Pipeline: {current_pipeline}", level=1)

        doc.add_heading(co["name"], level=2)

        # Meta-Zeile
        meta = doc.add_paragraph()
        mr = meta.add_run(
            f"Score {lead.get('overall_score', 0):.0%}  ·  "
            f"{co.get('industry', '')}  ·  {co.get('address', '')}"
        )
        mr.font.size = Pt(8)
        mr.font.color.rgb = RGBColor(120, 120, 120)

        # Ansprechpartner
        if contact_name:
            cp = doc.add_paragraph()
            cpr = cp.add_run(f"Ansprechpartner: {contact_name}")
            cpr.font.size = Pt(9)
            cpr.bold = True

        # Betreff
        ps = doc.add_paragraph()
        ps.add_run("Betreff: ").bold = True
        sr = ps.add_run(outreach["subject"])
        sr.italic = True
        sr.font.color.rgb = RGBColor(41, 65, 122)

        # Brieftext
        doc.add_paragraph(sal)
        doc.add_paragraph(intro)
        if outreach.get("hook"):
            doc.add_paragraph("Konkret zu Ihnen: " + outreach["hook"])
        doc.add_paragraph(CTA)
        doc.add_paragraph("Mit freundlichen Grüßen\nMarc Höppner · neogeo")
        doc.add_paragraph()

    # Quick Reference
    doc.add_page_break()
    doc.add_heading("Quick Reference — alle TIER-1-Betreffs", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Pipeline", "Firma", "Score", "Betreff"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for lead in tier1:
        row = table.add_row().cells
        row[0].text = lead.get("pipeline") or ""
        row[1].text = lead["company"]["name"]
        row[2].text = f"{lead.get('overall_score', 0):.0%}"
        row[3].text = lead["outreach"]["subject"]

    out = "iot_outreach_vorlagen.docx"
    doc.save(out)
    print(f"Word erstellt: {out} ({len(tier1)} Anschreiben)")


if __name__ == "__main__":
    main()
