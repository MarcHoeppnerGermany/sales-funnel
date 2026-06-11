#!/usr/bin/env python3
"""Excel-Export aller Leads — liest AUSSCHLIESSLICH aus leads.json.

Keine hardcodeten Daten: Kontakte, Tier, Pipeline und Outreach kommen aus
den Lead-Feldern (siehe sales_funnel/models/schemas.py).

Vorher Validierungs-Gate ausführen: python3 scripts/validate_leads.py
"""
import json
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="29417A", end_color="29417A", fill_type="solid")
TIER_FILLS = {
    "TIER1": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
    "TIER2": PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid"),
    "TIER3": PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid"),
    "SONDER": PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid"),
}
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def score_font(val):
    if val >= 0.7:
        return Font(name="Calibri", bold=True, color="2E8B57", size=10)
    if val >= 0.5:
        return Font(name="Calibri", bold=True, color="DAA520", size=10)
    return Font(name="Calibri", bold=True, color="B22222", size=10)


def style_header(ws, n_cols):
    for col in range(1, n_cols + 1):
        c = ws.cell(row=1, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
        c.border = BORDER


def auto_width(ws, min_w=8, max_w=45):
    for col_cells in ws.columns:
        lengths = [
            max(len(line) for line in str(c.value).split("\n"))
            for c in col_cells if c.value
        ]
        if lengths:
            w = min(max(max(lengths) + 2, min_w), max_w)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = w


def fmt_contacts(lead):
    parts = []
    for c in lead.get("contact_persons") or []:
        line = c["name"]
        if c.get("role"):
            line += f" | {c['role']}"
        if c.get("confidence"):
            line += f" | {c['confidence']}"
        parts.append(line)
    return "\n".join(parts)


def main():
    # Validierungs-Gate: kein Export auf kaputten Daten
    gate = subprocess.run(
        [sys.executable, "scripts/validate_leads.py"], capture_output=True, text=True
    )
    if gate.returncode != 0:
        print(gate.stdout)
        sys.exit("ABBRUCH: Validierungs-Gate fehlgeschlagen — erst Daten reparieren.")

    with open("leads.json") as f:
        leads = json.load(f)["leads"]
    leads.sort(key=lambda x: -(x.get("overall_score") or 0))

    wb = Workbook()

    # ---- Sheet 1: Alle Leads ----
    ws = wb.active
    ws.title = "Alle Leads"
    headers = [
        "Nr", "Pipeline", "Tier", "Firma", "Branche", "Standort", "Website",
        "Umsatz", "MA", "Overall", "Rating", "Gesundheit", "Affinität",
        "Size Fit", "Druck", "Hiring", "Topic-Hiring", "Ansprechpartner",
        "Outreach-Betreff",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, len(headers))

    for row, lead in enumerate(leads, 2):
        co = lead["company"]
        tr = lead.get("topic_relevance") or {}
        fin = lead.get("financials") or {}
        sf = fin.get("strategic_fit") or {}
        outreach = lead.get("outreach") or {}
        values = [
            row - 1, lead.get("pipeline"), lead.get("tier"), co["name"],
            co.get("industry"), co.get("address"), co.get("website"),
            fin.get("revenue_estimate"), fin.get("employee_count"),
            lead.get("overall_score"), fin.get("customer_rating"),
            (tr.get("company_health") or {}).get("activity_score"),
            (tr.get("topic_affinity") or {}).get("affinity_score"),
            sf.get("size_fit_score"), sf.get("pressure_score"),
            (fin.get("hiring_activity") or {}).get("general_hiring_score"),
            (fin.get("topic_hiring") or {}).get("topic_hiring_score"),
            fmt_contacts(lead), outreach.get("subject"),
        ]
        fill = TIER_FILLS.get(lead.get("tier") or "")
        for col, val in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.alignment = WRAP
            c.border = BORDER
            if fill:
                c.fill = fill
            if col in (10, 11) and isinstance(val, (int, float)):
                c.number_format = "0%"
                c.font = score_font(val)
            elif col in range(12, 18) and isinstance(val, (int, float)):
                c.number_format = "0%"

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(leads) + 1}"
    ws.freeze_panes = "A2"
    auto_width(ws)

    # ---- Sheet 2: je Pipeline ein Sheet ----
    pipelines = sorted({l.get("pipeline") for l in leads if l.get("pipeline")})
    for pipeline in pipelines:
        sub = [l for l in leads if l.get("pipeline") == pipeline]
        ws_p = wb.create_sheet(pipeline[:28])
        h2 = ["Nr", "Firma", "Overall", "Tier", "Ansprechpartner",
              "Outreach-Betreff", "Outreach-Aufhänger", "Strategie-Begründung"]
        for col, h in enumerate(h2, 1):
            ws_p.cell(row=1, column=col, value=h)
        style_header(ws_p, len(h2))
        for row, lead in enumerate(sub, 2):
            fin = lead.get("financials") or {}
            sf = fin.get("strategic_fit") or {}
            outreach = lead.get("outreach") or {}
            values = [
                row - 1, lead["company"]["name"], lead.get("overall_score"),
                lead.get("tier"), fmt_contacts(lead),
                outreach.get("subject"), outreach.get("hook"),
                sf.get("strategic_reasoning"),
            ]
            fill = TIER_FILLS.get(lead.get("tier") or "")
            for col, val in enumerate(values, 1):
                c = ws_p.cell(row=row, column=col, value=val)
                c.alignment = WRAP
                c.border = BORDER
                if fill:
                    c.fill = fill
                if col == 3 and isinstance(val, (int, float)):
                    c.number_format = "0%"
                    c.font = score_font(val)
        ws_p.freeze_panes = "A2"
        auto_width(ws_p, max_w=55)

    out = "iot_leads_komplett.xlsx"
    wb.save(out)
    print(f"Excel erstellt: {out} ({len(leads)} Leads, {1 + len(pipelines)} Sheets: "
          f"Alle + {', '.join(pipelines)})")


if __name__ == "__main__":
    main()
